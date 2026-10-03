"""Excel export helper (spec 54)."""

import io

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def build_workbook(sheets):
    """sheets: [(title, [columns], [[row values]])]."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    for title, columns, rows in sheets:
        sheet = workbook.create_sheet(title=title[:31])
        sheet.append(list(columns))
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            sheet.append(list(row))

        # Width from content, capped so one long description does not make the
        # sheet unusable.
        for index, column in enumerate(columns, start=1):
            longest = max(
                [len(str(column))] + [len(str(r[index - 1])) for r in rows[:200]
                                      if index - 1 < len(r)],
                default=10,
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 50)
        sheet.freeze_panes = "A2"
    return workbook


def workbook_response(workbook, filename):
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    response = HttpResponse(
        stream.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response
